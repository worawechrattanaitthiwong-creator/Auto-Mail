from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.processor import process_all, validate_upload_set


def make_transfer(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transfer Order Detail"
    ws.append(["Transfer number", "From warehouse", "To warehouse", "Transfer status", "Created date"])
    ws.append(["T1", "SCX_FFM", "STORE", "Deleted", datetime(2026, 8, 20, 8, 0)])
    ws.append(["T2", "STORE", "SCX_FFM", "Received", datetime(2026, 8, 20, 8, 0)])
    ws.append(["T3", "STORE", "ZZ_CNN", "Created", datetime(2026, 8, 20, 8, 0)])
    ws.append(["T4", "STORE", "STORE", "Shipped", datetime(2025, 8, 20, 8, 0)])
    wb.save(path)


def make_purchase(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Order Detail"
    ws.append(["PurchaseCreatedDateTime", "INVENTLOCATIONID", "PURCHID"])
    for location in ["ZZ_CNN", "SCX_WH1", "SCX_FFM", "SCX_XD"]:
        ws.append([datetime(2026, 8, 20, 8, 0), location, location])
    ws.append([datetime(2025, 8, 20, 8, 0), "SCX_XD", "OLD"])
    wb.save(path)


def make_diff(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["DOCNO", "CREATEDDATETIME"])
    ws.append(["D1", datetime(2026, 8, 20, 8, 0)])
    wb.save(path)


def count_data_rows(path: Path) -> int:
    wb = load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    count = max(sum(1 for _ in ws.iter_rows(values_only=True)) - 1, 0)
    wb.close()
    return count


def test_validate_and_process(tmp_path: Path) -> None:
    names = [
        "TransferOrder_2026082008.xlsx",
        "PurchaseOrder_2026082008.xlsx",
        "TransferOrderDiff_2026082008.xlsx",
    ]
    _, report_date = validate_upload_set(names)
    assert report_date == date(2026, 8, 20)

    transfer = tmp_path / names[0]
    purchase = tmp_path / names[1]
    diff = tmp_path / names[2]
    make_transfer(transfer)
    make_purchase(purchase)
    make_diff(diff)

    outputs = process_all(transfer, purchase, diff, tmp_path / "out", report_date)
    assert len(outputs) == 12
    assert outputs["transfer_order_diff"].name == "TransferOrderDiff_ 20-08-2026 Time 09.00.xlsx"
    assert count_data_rows(outputs["transfer_status_deleted"]) == 1
    assert count_data_rows(outputs["transfer_status_received"]) == 1
    assert count_data_rows(outputs["transfer_status_created"]) == 1
    assert count_data_rows(outputs["transfer_status_shipped"]) == 0
    assert count_data_rows(outputs["transfer_scx_ffm_out"]) == 1
    assert count_data_rows(outputs["transfer_scx_ffm_in"]) == 1
    assert count_data_rows(outputs["transfer_zz_cnn_in"]) == 1
    assert count_data_rows(outputs["purchase_prt_zz_cnn"]) == 1
    assert count_data_rows(outputs["purchase_rc_scx_wh1"]) == 1
    assert count_data_rows(outputs["purchase_rc_scx_ffm"]) == 1
    assert count_data_rows(outputs["purchase_rc_scx_xd"]) == 1
