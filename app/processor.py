from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook

ProgressCallback = Callable[[int, str], None]

TRANSFER_STATUS_VALUES = ("Received", "Created", "Deleted", "Shipped")
PURCHASE_LOCATION_PREFIXES = {
    "ZZ_CNN": "PRT.ZZ_CNN",
    "SCX_WH1": "Rc.SCX_WH1",
    "SCX_FFM": "Rc.SCX_FFM",
    "SCX_XD": "Rc.SCX_XD",
}

FILE_PATTERNS = {
    "transfer_order_diff": re.compile(r"^TransferOrderDiff_(\d{8})\d*\.xlsx$", re.I),
    "transfer_order": re.compile(r"^TransferOrder_(\d{8})\d*\.xlsx$", re.I),
    "purchase_order": re.compile(r"^PurchaseOrder_(\d{8})\d*\.xlsx$", re.I),
}


class ProcessingError(RuntimeError):
    pass


@dataclass(frozen=True)
class ClassifiedFile:
    kind: str
    filename: str
    report_date: date


def classify_filename(filename: str) -> ClassifiedFile:
    for kind, pattern in FILE_PATTERNS.items():
        match = pattern.match(filename)
        if match:
            report_date = datetime.strptime(match.group(1), "%Y%m%d").date()
            return ClassifiedFile(kind=kind, filename=filename, report_date=report_date)
    raise ProcessingError(
        f"ชื่อไฟล์ไม่ตรงรูปแบบที่รองรับ: {filename}. "
        "ต้องเป็น TransferOrder_YYYYMMDD..., PurchaseOrder_YYYYMMDD... "
        "หรือ TransferOrderDiff_YYYYMMDD..."
    )


def validate_upload_set(filenames: Iterable[str]) -> tuple[dict[str, ClassifiedFile], date]:
    classified: dict[str, ClassifiedFile] = {}
    for filename in filenames:
        item = classify_filename(filename)
        if item.kind in classified:
            raise ProcessingError(f"พบไฟล์ชนิด {item.kind} ซ้ำมากกว่า 1 ไฟล์")
        classified[item.kind] = item

    required = set(FILE_PATTERNS)
    missing = required - set(classified)
    if missing:
        raise ProcessingError("ไฟล์ไม่ครบ 3 ชนิด: " + ", ".join(sorted(missing)))

    dates = {item.report_date for item in classified.values()}
    if len(dates) != 1:
        pretty = ", ".join(sorted(d.strftime("%d-%m-%Y") for d in dates))
        raise ProcessingError(f"วันที่ในชื่อไฟล์ทั้ง 3 ไฟล์ไม่ตรงกัน: {pretty}")

    return classified, dates.pop()


def _date_text(report_date: date) -> str:
    return report_date.strftime("%d-%m-%Y")


def _year_matches(value: object, year: int) -> bool:
    if isinstance(value, (datetime, date)):
        return value.year == year
    if isinstance(value, str):
        text = value.strip()
        # Common source formats are YYYY-MM-DD[ HH:MM:SS], DD/MM/YYYY and MM/DD/YYYY.
        # Avoid datetime.strptime in the ~1M-row hot loop.
        if len(text) >= 4 and text[:4].isdigit():
            return int(text[:4]) == year
        if len(text) >= 10 and text[6:10].isdigit() and text[2] in "/-" and text[5] in "/-":
            return int(text[6:10]) == year
    return False


def _require_headers(header: tuple[object, ...], names: Iterable[str], source_name: str) -> dict[str, int]:
    lookup = {str(value).strip(): index for index, value in enumerate(header) if value is not None}
    missing = [name for name in names if name not in lookup]
    if missing:
        raise ProcessingError(f"{source_name} ขาดคอลัมน์ที่จำเป็น: {', '.join(missing)}")
    return lookup


def _new_output_book(header: tuple[object, ...]) -> tuple[Workbook, object]:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Sheet1")
    sheet.append(header)
    return workbook, sheet


def _save_books(
    books: dict[str, tuple[Workbook, object]],
    paths: dict[str, Path],
    progress: ProgressCallback | None = None,
    *,
    start_pct: int,
    end_pct: int,
    label: str,
) -> None:
    total = max(len(books), 1)
    span = max(end_pct - start_pct, 0)
    for index, (key, (workbook, _)) in enumerate(books.items(), start=1):
        workbook.save(paths[key])
        if progress:
            pct = start_pct + int(index / total * span)
            progress(pct, f"กำลังบันทึก {label}: {index}/{total} ไฟล์")


def process_transfer_order(
    source: Path,
    output_dir: Path,
    report_date: date,
    progress: ProgressCallback | None = None,
) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = tuple(next(rows))
    except StopIteration as exc:
        workbook.close()
        raise ProcessingError("TransferOrder ไม่มีข้อมูล") from exc

    columns = _require_headers(
        header,
        ("From warehouse", "To warehouse", "Transfer status", "Created date"),
        "TransferOrder",
    )
    report_year = report_date.year
    stamp = f"{_date_text(report_date)} 19.00"

    paths: dict[str, Path] = {
        f"transfer_status_{status.lower()}": output_dir
        / f"Tranfer order lines Status {status} {stamp}.xlsx"
        for status in TRANSFER_STATUS_VALUES
    }
    paths.update(
        {
            "transfer_scx_ffm_out": output_dir / f"Dynamics SCX_FFM Out {stamp}.xlsx",
            "transfer_scx_ffm_in": output_dir / f"Dynamics SCX_FFM IN {stamp}.xlsx",
            "transfer_zz_cnn_in": output_dir / f"Dynamics ZZ_CNN IN {stamp}.xlsx",
        }
    )

    books = {key: _new_output_book(header) for key in paths}
    total = max((sheet.max_row or 1) - 1, 1)

    try:
        for row_number, row in enumerate(rows, start=2):
            if not _year_matches(row[columns["Created date"]], report_year):
                continue

            status = str(row[columns["Transfer status"]] or "").strip()
            status_key = f"transfer_status_{status.lower()}"
            if status_key in books:
                books[status_key][1].append(row)

            if str(row[columns["From warehouse"]] or "").strip() == "SCX_FFM":
                books["transfer_scx_ffm_out"][1].append(row)

            to_warehouse = str(row[columns["To warehouse"]] or "").strip()
            if to_warehouse == "SCX_FFM":
                books["transfer_scx_ffm_in"][1].append(row)
            if to_warehouse == "ZZ_CNN":
                books["transfer_zz_cnn_in"][1].append(row)

            if progress and row_number % 10000 == 0:
                pct = min(38, int((row_number - 1) / total * 38))
                progress(pct, f"กำลังแยก TransferOrder: {row_number - 1:,}/{total:,} แถว")
    finally:
        workbook.close()

    _save_books(
        books,
        paths,
        progress,
        start_pct=38,
        end_pct=45,
        label="TransferOrder",
    )
    if progress:
        progress(45, "แยก TransferOrder เป็น 7 ไฟล์เรียบร้อย")
    return paths


def process_purchase_order(
    source: Path,
    output_dir: Path,
    report_date: date,
    progress: ProgressCallback | None = None,
) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = tuple(next(rows))
    except StopIteration as exc:
        workbook.close()
        raise ProcessingError("PurchaseOrder ไม่มีข้อมูล") from exc

    columns = _require_headers(header, ("PurchaseCreatedDateTime", "INVENTLOCATIONID"), "PurchaseOrder")
    report_year = report_date.year
    stamp = f"{_date_text(report_date)} 19.00"

    paths = {
        f"purchase_{prefix.lower().replace('.', '_')}": output_dir / f"{prefix} {stamp}.xlsx"
        for prefix in PURCHASE_LOCATION_PREFIXES.values()
    }
    location_to_key = {
        location: f"purchase_{prefix.lower().replace('.', '_')}"
        for location, prefix in PURCHASE_LOCATION_PREFIXES.items()
    }

    books = {key: _new_output_book(header) for key in paths}
    total = max((sheet.max_row or 1) - 1, 1)

    try:
        for row_number, row in enumerate(rows, start=2):
            if not _year_matches(row[columns["PurchaseCreatedDateTime"]], report_year):
                continue
            location = str(row[columns["INVENTLOCATIONID"]] or "").strip()
            key = location_to_key.get(location)
            if key:
                books[key][1].append(row)

            if progress and row_number % 10000 == 0:
                pct = 45 + min(34, int((row_number - 1) / total * 34))
                progress(pct, f"กำลังแยก PurchaseOrder: {row_number - 1:,}/{total:,} แถว")
    finally:
        workbook.close()

    _save_books(
        books,
        paths,
        progress,
        start_pct=79,
        end_pct=85,
        label="PurchaseOrder",
    )
    if progress:
        progress(85, "แยก PurchaseOrder เป็น 4 ไฟล์เรียบร้อย")
    return paths


def rename_transfer_order_diff(source: Path, output_dir: Path, report_date: date) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / f"TransferOrderDiff_ {_date_text(report_date)} Time 19.00.xlsx"
    shutil.copy2(source, target)
    return {"transfer_order_diff": target}


def process_all(
    transfer_order: Path,
    purchase_order: Path,
    transfer_order_diff: Path,
    output_dir: Path,
    report_date: date,
    progress: ProgressCallback | None = None,
) -> dict[str, Path]:
    outputs: dict[str, Path] = {}
    if progress:
        progress(1, "เริ่มประมวลผลไฟล์")
    outputs.update(process_transfer_order(transfer_order, output_dir, report_date, progress))
    outputs.update(process_purchase_order(purchase_order, output_dir, report_date, progress))
    outputs.update(rename_transfer_order_diff(transfer_order_diff, output_dir, report_date))
    if progress:
        progress(87, "เตรียมไฟล์ครบ 12 ไฟล์แล้ว")
    return outputs
